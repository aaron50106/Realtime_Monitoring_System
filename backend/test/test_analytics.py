from datetime import datetime, timedelta, timezone
from io import BytesIO
import uuid

from openpyxl import load_workbook
import pytest

from conftest import TEST_PREFIX


TAIPEI_WINDOW = {
    "start_time": "2026-10-21T00:00:00+08:00",
    "end_time": "2026-10-21T23:59:59+08:00",
}
UTC_WINDOW = {
    "start_time": "2026-10-20T16:00:00Z",
    "end_time": "2026-10-21T15:59:59Z",
}
ENDPOINTS = ["summary", "by-category", "trend", "export/excel"]
EXCEL_HEADERS = (
    "id", "title", "value", "category", "timestamp", "is_alert",
    "creator_id", "timestamp_taipei",
)


@pytest.fixture
def analytics_data(integration_client, create_user):
    user = create_user(role="admin")
    login = integration_client.post("/api/v1/auth/login", json={
        "username": user["username"], "password": user["password"],
    })
    assert login.status_code == 200
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    category = f"{TEST_PREFIX}tz_{uuid.uuid4().hex[:8]}"
    samples = [
        ("2026-10-20T23:59:59+08:00", 5),
        ("2026-10-21T00:00:00+08:00", 10),
        ("2026-10-21T00:00:59+08:00", 20),
        ("2026-10-21T00:01:00+08:00", 30),
        ("2026-10-21T07:59:59+08:00", 40),
        ("2026-10-21T08:00:00+08:00", 50),
        ("2026-10-21T23:59:59+08:00", 60),
        ("2026-10-22T00:00:00+08:00", 70),
    ]
    imported = integration_client.post(
        "/api/v1/data-records/import/json",
        headers=headers,
        json={"records": [
            {
                "title": f"{category}-{index}", "value": value,
                "category": category, "timestamp": timestamp,
            }
            for index, (timestamp, value) in enumerate(samples)
        ]},
    )
    assert imported.status_code == 201
    assert imported.json() == {"inserted_count": 8}
    return {"headers": headers, "category": category, "user_id": user["id"]}


def _excel_rows(response):
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        return list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()


@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_analytics_equivalent_offsets(integration_client, analytics_data, endpoint):
    windows = [
        TAIPEI_WINDOW,
        UTC_WINDOW,
        {key: value.removesuffix("Z") for key, value in UTC_WINDOW.items()},
        {
            "start_time": "2026-10-20T12:00:00-04:00",
            "end_time": "2026-10-21T11:59:59-04:00",
        },
        {"start_time": TAIPEI_WINDOW["start_time"], "end_time": UTC_WINDOW["end_time"]},
    ]
    results = []
    for window in windows:
        response = integration_client.get(
            f"/api/v1/analytics/{endpoint}",
            headers=analytics_data["headers"],
            params={"category": analytics_data["category"], **window},
        )
        assert response.status_code == 200, response.text
        results.append(_excel_rows(response) if endpoint == "export/excel" else response.json())
    assert all(result == results[0] for result in results)
    result = results[0]
    if endpoint == "summary":
        assert result == {"total_count": 6, "avg_value": 35, "min_value": 10, "max_value": 60}
    elif endpoint == "by-category":
        assert result["total_categories"] == 1
        assert result["items"][0] == {
            "category": analytics_data["category"], "count": 6,
            "avg_value": 35, "min_value": 10, "max_value": 60, "sum_value": 210,
        }
    elif endpoint == "trend":
        assert [point["count"] for point in result["points"]] == [3, 1, 1, 1]
    else:
        assert result[0] == EXCEL_HEADERS
        assert len(result) == 7
        assert result[1][4] == "2026-10-20T16:00:00+00:00"
        assert result[1][7] == "2026-10-21T00:00:00+08:00"
        assert result[-1][4] == "2026-10-21T15:59:59+00:00"
        assert result[-1][7] == "2026-10-21T23:59:59+08:00"


@pytest.mark.parametrize(
    "granularity, expected_buckets, expected_counts",
    [
        ("minute", [
            "2026-10-20T15:59:00Z", "2026-10-20T16:00:00Z",
            "2026-10-20T16:01:00Z", "2026-10-20T23:59:00Z",
            "2026-10-21T00:00:00Z", "2026-10-21T15:59:00Z",
            "2026-10-21T16:00:00Z",
        ], [1, 2, 1, 1, 1, 1, 1]),
        ("hour", [
            "2026-10-20T15:00:00Z", "2026-10-20T16:00:00Z",
            "2026-10-20T23:00:00Z", "2026-10-21T00:00:00Z",
            "2026-10-21T15:00:00Z", "2026-10-21T16:00:00Z",
        ], [1, 3, 1, 1, 1, 1]),
        ("day", [
            "2026-10-19T16:00:00Z", "2026-10-20T16:00:00Z", "2026-10-21T16:00:00Z",
        ], [1, 6, 1]),
    ],
)
def test_trend_uses_taiwan_calendar(
    integration_client, analytics_data, granularity, expected_buckets, expected_counts,
):
    response = integration_client.get(
        "/api/v1/analytics/trend",
        headers=analytics_data["headers"],
        params={"category": analytics_data["category"], "granularity": granularity},
    )
    assert response.status_code == 200
    points = response.json()["points"]
    buckets = [datetime.fromisoformat(point["bucket"]) for point in points]
    assert buckets == [datetime.fromisoformat(value) for value in expected_buckets]
    assert all(bucket.utcoffset() == timedelta(0) for bucket in buckets)
    assert [point["count"] for point in points] == expected_counts
    if granularity == "day":
        taipei = timezone(timedelta(hours=8))
        assert [bucket.astimezone(taipei).hour for bucket in buckets] == [0, 0, 0]
        assert points[1]["avg_value"] == 35


@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_analytics_rejects_reversed_offset_window(integration_client, analytics_data, endpoint):
    response = integration_client.get(
        f"/api/v1/analytics/{endpoint}",
        headers=analytics_data["headers"],
        params={
            "start_time": "2026-10-21T00:00:00+08:00",
            "end_time": "2026-10-20T15:59:59Z",
        },
    )
    assert response.status_code == 422
    assert response.json() == {"detail": "start_time must be less than or equal to end_time"}


@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_analytics_empty_time_window(integration_client, analytics_data, endpoint):
    response = integration_client.get(
        f"/api/v1/analytics/{endpoint}",
        headers=analytics_data["headers"],
        params={
            "category": analytics_data["category"],
            "start_time": "2026-10-23T00:00:00+08:00",
            "end_time": "2026-10-23T00:00:01+08:00",
        },
    )
    assert response.status_code == 200
    if endpoint == "summary":
        assert response.json() == {
            "total_count": 0, "avg_value": None, "min_value": None, "max_value": None,
        }
    elif endpoint == "by-category":
        assert response.json() == {"items": [], "total_categories": 0}
    elif endpoint == "trend":
        assert response.json() == {"granularity": "hour", "points": []}
    else:
        assert _excel_rows(response) == [EXCEL_HEADERS]


def test_admin_created_at_has_explicit_utc(integration_client, analytics_data):
    response = integration_client.get(
        f"/api/v1/admin/users/{analytics_data['user_id']}",
        headers=analytics_data["headers"],
    )
    assert response.status_code == 200
    created_at = datetime.fromisoformat(response.json()["created_at"])
    assert created_at.utcoffset() == timedelta(0)
    assert abs((datetime.now(timezone.utc) - created_at).total_seconds()) < 60
